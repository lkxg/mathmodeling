"""Build a fully editable Chinese report from checked Q1 artifacts."""
from __future__ import annotations
import json
from pathlib import Path
import hashlib
import numpy as np
import pandas as pd
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from ..data import default_config, load_samples, output_dir
from ....common import read_json, file_hash
from . import figures, equations

PROJECT=Path(__file__).resolve().parents[4]
OUT=PROJECT/"outputs/q1/documentation"
NAME="问题1_方法流程验证与结果报告"
INK="173644";TEAL="147D86";GRAY="526A76";PALE="F0F6F7"
BODY="Noto Serif CJK SC";SANS="Noto Sans CJK SC"


def font(run,size=10.5,bold=False,color=INK,family=BODY):
    run.font.name=family;run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor.from_string(color)
    rf=run._element.get_or_add_rPr().get_or_add_rFonts()
    for k in ("ascii","hAnsi","eastAsia","cs"):rf.set(qn("w:"+k),family)


class Report:
    def __init__(self):
        self.doc=Document();s=self.doc.sections[0]
        s.page_width=Cm(21);s.page_height=Cm(29.7);s.top_margin=Cm(1.8);s.bottom_margin=Cm(1.7)
        s.left_margin=Cm(2);s.right_margin=Cm(2);s.header_distance=Cm(.75);s.footer_distance=Cm(.75)
        s.different_first_page_header_footer=True
        for style in ("Normal","Body Text","Caption"):
            st=self.doc.styles[style];st.font.name=BODY;st.font.size=Pt(10.5)
            st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"),BODY)
            st.paragraph_format.line_spacing=1.22;st.paragraph_format.space_after=Pt(6)
        for level,size in [(1,19),(2,12.5),(3,11)]:
            st=self.doc.styles[f"Heading {level}"];st.font.name=SANS;st.font.size=Pt(size);st.font.bold=True
            st.font.color.rgb=RGBColor.from_string(TEAL);st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"),SANS)
            st.paragraph_format.space_before=Pt(9 if level>1 else 0);st.paragraph_format.space_after=Pt(8)
        h=s.header.paragraphs[0];h.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        font(h.add_run("第一问｜多模态特征与可追溯时序对齐"),8.5,color=GRAY,family=SANS)
        f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
        font(f.add_run("第一问方法与结果报告  ·  "),8,color=GRAY)
        field=OxmlElement("w:fldSimple");field.set(qn("w:instr"),"PAGE");f._p.append(field)
        prop=self.doc.core_properties;prop.title="第一问：多模态情感特征提取与时序对齐";prop.subject="方法、流程、验证与结果"
        prop.author="";prop.last_modified_by="";prop.keywords="多模态;时序对齐;可追溯;特征提取";prop.comments="由本地实现与已核验实验产物生成；不包含人工词边界真值。"
        self.toc=[];self.figures=0;self.tables=0;self.eq=0
        self.page_map=read_json(OUT/"page_index.json") if (OUT/"page_index.json").exists() else {}
        self.toc_table=None

    def p(self,text="",size=10.5,bold=False,color=INK,after=6,align=None):
        p=self.doc.add_paragraph();p.paragraph_format.space_after=Pt(after)
        if align is not None:p.alignment=align
        font(p.add_run(text),size,bold,color);return p

    def page(self,title,subtitle="",toc=True):
        self.doc.add_page_break();p=self.doc.add_paragraph(title,"Heading 1")
        if toc:
            self.toc.append(title);number=len(self.toc)
            start=OxmlElement("w:bookmarkStart");start.set(qn("w:id"),str(number));start.set(qn("w:name"),f"section_{number}")
            end=OxmlElement("w:bookmarkEnd");end.set(qn("w:id"),str(number));p._p.insert(0,start);p._p.append(end)
        if subtitle:self.p(subtitle,9.2,color=GRAY,after=10)
        return p

    def h(self,text):self.doc.add_paragraph(text,"Heading 2")

    def note(self,text):
        p=self.p(text,9.3,color=TEAL,after=8)
        pp=p._p.get_or_add_pPr();shd=OxmlElement("w:shd");shd.set(qn("w:fill"),PALE);pp.append(shd)
        return p

    def equation(self,text):
        self.eq+=1;p=self.doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after=Pt(8)
        p._p.append(equations.formula(self.eq))
        font(p.add_run(f"    （{self.eq}）"),9,color=GRAY)

    def table(self,headers,rows,widths=None,title=None,size=9):
        self.tables+=1
        if title:self.p(f"表{self.tables}  {title}",9.5,bold=True,color=TEAL,after=5).paragraph_format.keep_with_next=True
        t=self.doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
        if widths is None:widths=[17/len(headers)]*len(headers)
        for col,w in zip(t.columns,widths):col.width=Cm(w)
        for ri,values in enumerate([headers]+list(rows)):
            row=t.rows[0] if ri==0 else t.add_row()
            trpr=row._tr.get_or_add_trPr();cant=OxmlElement("w:cantSplit");trpr.append(cant)
            if ri==0:rep=OxmlElement("w:tblHeader");trpr.append(rep)
            for j,value in enumerate(values):
                c=row.cells[j];c.width=Cm(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p=c.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1.10
                font(p.add_run(str(value)),size,ri==0,"FFFFFF" if ri==0 else INK,family=SANS if ri==0 else BODY)
                tcpr=c._tc.get_or_add_tcPr();shade=OxmlElement("w:shd");shade.set(qn("w:fill"),TEAL if ri==0 else ("F1F5F6" if ri%2 else "FFFFFF"));tcpr.append(shade)
                mar=OxmlElement("w:tcMar")
                for k,v in [("top",65),("bottom",65),("left",85),("right",85)]:
                    e=OxmlElement("w:"+k);e.set(qn("w:w"),str(v));e.set(qn("w:type"),"dxa");mar.append(e)
                tcpr.append(mar)
        self.p("",size=1,after=3)
        return t

    def image(self,path,caption,width=16.7):
        self.figures+=1;p=self.doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(3);p.paragraph_format.keep_with_next=True
        p.add_run().add_picture(str(path),width=Cm(width))
        self.p(f"图{self.figures}  {caption}",9,color=GRAY,after=8,align=WD_ALIGN_PARAGRAPH.CENTER)

    def code(self,text):
        p=self.p(text,8.2,color=INK,after=7);p.paragraph_format.line_spacing=1.05
        for r in p.runs:font(r,8.2,family="DejaVu Sans Mono")
        pp=p._p.get_or_add_pPr();sh=OxmlElement("w:shd");sh.set(qn("w:fill"),"F2F5F6");pp.append(sh)

    def link(self,label,url):
        p=self.doc.add_paragraph();p.paragraph_format.space_after=Pt(5)
        rid=self.doc.part.relate_to(url,RT.HYPERLINK,is_external=True)
        h=OxmlElement("w:hyperlink");h.set(qn("r:id"),rid);r=OxmlElement("w:r");pr=OxmlElement("w:rPr")
        color=OxmlElement("w:color");color.set(qn("w:val"),TEAL);pr.append(color);r.append(pr)
        text=OxmlElement("w:t");text.text=label;r.append(text);h.append(r);p._p.append(h)
        return p

    def fill_toc(self):
        for number,title in enumerate(self.toc,start=1):
            row=self.toc_table.add_row()
            for i,val in enumerate([title,str(self.page_map.get(title,"—"))]):
                p=row.cells[i].paragraphs[0];p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.03
                r=p.add_run(val);font(r,9.2,color=INK)
                if i==0:
                    link=OxmlElement("w:hyperlink");link.set(qn("w:anchor"),f"section_{number}")
                    p._p.remove(r._r);link.append(r._r);p._p.append(link)
                if i==1:p.alignment=WD_ALIGN_PARAGRAPH.RIGHT


def f4(x):return "—" if x is None or not np.isfinite(x) else f"{x:.4f}"


def run():
    OUT.mkdir(parents=True,exist_ok=True);cfg=default_config();root=Path(cfg["output_dir"]);study=output_dir(cfg,"")
    samples=load_samples(cfg);summary=pd.read_csv(root/"summary.csv");probe=read_json(root/"analysis/probe_v2/report.json")
    reports={n:read_json(study/n/"report.json") for n in ("validation","anchors","sampling","soft","review")}
    validation=reports["validation"];vad=read_json(root/"analysis/alignment_vad.json");runconfig=read_json(root/"run_config.json")
    for r in reports.values():
        assert all(file_hash(p)==h for p,h in r["provenance"]["code_files"].items()),"Stale experiment report"
    assert read_json(study/"verification.json")["passed"]
    figures_dir=OUT/"figures"
    figs=figures.build(figures_dir,samples,summary,probe,reports,study)
    D=Report();v=validation;totalwords=int(summary.words.sum());validwords=sum(int(sum(w["time_valid"] for w in s["words"])) for s in samples)
    positions=int(summary.positions.sum());nativeframes=sum(len(s["vision"]["quality"]) for s in samples);detected=sum(int((s["vision"]["quality"]>0).sum()) for s in samples)
    D.p("E 题  ·  第一问技术报告",13,bold=True,color=TEAL,after=30)
    D.p("多模态情感特征提取\n与时序对齐",29,bold=True,color=INK,after=20)
    D.p("方法 · 流程 · 验证 · 结果",18,color=TEAL,after=24)
    D.p("基于附件1的100条原始视频及已有英文转写",12,after=5)
    D.p("版本 1.0    |    结果快照：2026年9月24日",10,color=GRAY,after=34)
    D.table(["完整样本","时间轴覆盖","自动测试"],[["100 / 100",f"{100*v['mapped_seconds']/v['clip_seconds']:.4f}%","35 项通过"]],widths=[5.67,5.66,5.67],title="报告核心指标",size=14)
    D.p("可追溯的共同时间坐标",16,bold=True,color=TEAL,after=10)
    D.p("保留各模态原生序列，以词区间与未分配间隙组织输出；通过时间重叠与有效质量加权，将文本、声学与视觉观测对应回同一段原始素材。",12,after=22)
    D.note("报告范围：已实现的主流程、已完成的辅助实验与当前输出。人工词边界尚未标注；已知8帧视觉误报的复核掩码尚未合入主特征。")
    D.p("所有数值均来自本地运行产物。文中区分形式有效、观测可用与人工真实性验证，不把覆盖率解释为词边界准确率。",10,color=GRAY,after=10)
    D.page("阅读导航",toc=False)
    D.p("正文按“数据输入 → 特征定义 → 时间模型 → 输出规范 → 验证结果 → 局限与复现”组织。附录提供全部100条样本、原始声学维度及模型版本信息。",10)
    D.toc_table=D.doc.add_table(rows=0,cols=2);D.toc_table.autofit=False;D.toc_table.columns[0].width=Cm(15.5);D.toc_table.columns[1].width=Cm(1.5)

    D.page("1  任务边界与数据基础","第一问解决特征提取与时序组织；辅助情感预测用于诊断特征设计。")
    D.p("题目允许使用开源工具与预训练模型完成基础处理，要求原始样本覆盖完整、时序组织可核验、方法合理且可复现，并展示至少一个典型样本。本报告以当前代码和输出为依据，重点说明每个特征来自哪个原始位置及其有效性。[D1]",10.5)
    D.table(["任务要求","本方案的对应产物"],[
        ["完整覆盖与一一对应","100条清单；样本ID、video_id、clip_id、标签与视频SHA256"],
        ["时序组织可核验","变长序列、有效长度、填充掩码、原始索引和CSR映射"],
        ["方法与复现","固定模型修订号、配置、运行日志、依赖版本与核验命令"],
        ["典型样本展示","原文、语音时段、视频帧及特征曲线的同轴展示"]],[4.0,13.0],title="题目要求与实现对应关系")
    D.p(f"数据共100条片段，来自37个原视频；总时长{v['clip_seconds']:.3f}秒，单片段范围{summary.duration.min():.3f}—{summary.duration.max():.3f}秒。共有{totalwords}个原文词，最终组织为{positions}个位置，其中{positions-totalwords}个为未分配区间。",10.5)
    D.image(figs["data"],"附件1的时长与标签分布；标签仅用于后续辅助验证。")
    D.note("附件2是问题2/3的训练与验证数据。第一问生成的768/818/56维空间，不能仅凭维度或序列长度与附件2直接互换；本报告不包含附件3/4专项测试结果。")

    D.page("2  总体流程与设计原则","原生时间轴保留 → 共同时间身份 → 可解释聚合 → 特征与来源同步导出。")
    D.image(figs["architecture"],"已实现主流程及独立保存的补充实验。",width=16.0)
    D.p("主流程按媒体解析、强制对齐、文本编码、语音编码、视觉编码、聚合六个阶段执行。媒体解析是共同上游；三模态提取结果先保留原生时间结构，再生成统一输出位置。阶段模型按需加载，缓存通过输入与产物哈希校验后复用。",10.2)
    D.note("主体流程只使用冻结预训练模型提取特征；探针和小型软对齐网络属于独立辅助实验。主输出不依赖这些辅助分类器的预测。")

    D.page("3  文本：保留原文并补充词级时间","已有Transcript → Qwen3强制对齐 → 原文字串映射 → ModernBERT上下文词表示。")
    D.p("文本内容直接采用题目提供的英文转写。强制对齐器接收音频和该转写，补充词级起止时间；不使用自由ASR结果替换原文。采用Qwen3-ForcedAligner的固定权重版本，推理输出保留原始音频时间与共同坐标时间。[R1]",10.5)
    D.p("ModernBERT读取完整原始文本，通过token字符offset与原文词字符跨度建立对应，排除特殊token；一个词对应多个子词时取其上下文表示均值，得到768维词向量。文本模型冻结，不在附件1上微调。[R2]",10.5)
    D.equation("xᵀᵢ = (1 / |Bᵢ|) ∑ⱼ∈Bᵢ hⱼ ,     xᵀᵢ ∈ ℝ⁷⁶⁸")
    D.p("其中Bᵢ是与第i个原文词字符区间有交集的有效子词集合，hⱼ为上下文隐状态。原文映射失败时直接报错，不用模糊匹配补造词身份。",10)
    D.table(["情况","文本处理","时间/音视频处理"],[
        ["正常词与时间","保留词向量和原文跨度","按有效词区间聚合音视频"],
        ["零时长、越界或重叠","词语义仍保留","time_valid=false，不强配音视频"],
        ["整段数字静音","给定转写仍保留","全部词时间无效；音频不可用"],
        ["原文字符映射不完整","明确失败并记录","不编造索引或时间"]],[4,5.3,7.7],title="语义有效与时间有效分别管理")
    s=samples[0]
    D.h("实际例子：一个词可以有语义、没有可靠时间")
    D.p(s["raw_text"],11,color=TEAL)
    D.p("此样本中“the”的原文跨度为[70,73)，自动时间为[4.48,4.48]秒，时长为零。该词保留文本表示，但其time_valid_mask为false。这个处理避免了用看似完整的音视频向量掩盖时间失败。",10.5)

    D.page("4  语音：学习表示与声学描述量","16 kHz单声道输入；分别保留emotion2vec与openSMILE的原生时间窗。")
    D.p("音轨按原始解码起点对齐到共同坐标，转换为16 kHz单声道浮点波形。音频采样下标与时间的换算使用实际音轨偏移；检测到不连续PTS时拒绝将其当作连续信号拼接。",10.5)
    D.table(["分量","原生内容与时间组织","导出维度"],[
        ["emotion2vec","冻结模型的帧级情感相关表示；卷积锚点窗口25 ms、步长20 ms","加权均值768"],
        ["eGeMAPSv02 LLD","响度、F0、谱形、MFCC、jitter/shimmer、谐噪比及共振峰等25项；使用工具返回时间窗","均值25＋标准差25"],
        ["最终语音表示","两条源序列分别聚合至同一目标区间后拼接","768＋25＋25＝818"]],[3.4,10.5,3.1],title="语音特征的来源与维度")
    D.equation("xᴬᵢ = [ μᵉᵢ ; μˡᵢ ; σˡᵢ ] ∈ ℝ⁸¹⁸")
    D.p("emotion2vec提供学习到的语音表示，LLD补充可解释的声学变化量；两类特征的用途不同。[R3,R4] 25维标准差是同一目标区间内源帧的变化，不是把多个词向量之间的标准差拿来替代，也不是模型预测不确定度。",10.5)
    D.h("时间解释与静音处理")
    D.p("25 ms/20 ms描述的是emotion2vec前端卷积的名义时间锚点。其Transformer隐状态仍包含上下文，因此不能声称每个768维向量只观察了独立的25 ms音频。",10.5)
    D.p("整段波形峰值≤10⁻⁴时，按数字静音处理：emotion2vec与LLD整段不可用，所有词时间无效，给定文本保留。该规则只针对整段静音，不等同于已经完成所有局部静音或背景噪声识别。",10.5)
    D.note("音频模态掩码为两个声学分量掩码的逻辑或。下游仍需查看component_mask，区分768维表示与25＋25维描述量是否分别可用。")

    D.page("5  视觉：面部观测、跟踪与质量记录","实际原帧PTS采样；单帧28维，再按目标时间区间导出均值与标准差。")
    D.p("视觉使用OpenFace 3.0相关检测与预测模块。[R5] 主流程按原始PTS的最小时间间隔近似10 Hz取帧，记录原帧索引及实际PTS。图像长边最多960像素；检测阈值0.8。先选最大人脸，后续优先选择与前一帧边框IoU最大的脸；IoU低于0.2时重选并记录跟踪重置。",10.5)
    D.table(["单帧分量","维度","含义与限制"],[
        ["表情logits","8","模型原始表情输出；未直接映射为赛题情感极性"],
        ["视线输出","2","按模型返回坐标保存"],
        ["AU原始输出","8","保留原始索引，不擅自标注为具体FACS编号"],
        ["五点人脸地标","10","相对检测边框归一化的(x,y)坐标"],
        ["单帧合计 / 区间导出","28 / 56","区间均值28＋区间标准差28"]],[4.0,2.3,10.7],title="视觉28维源向量与56维输出")
    D.equation("xⱽᵢ = [ μᵛᵢ ; σᵛᵢ ] ∈ ℝ⁵⁶")
    D.p("未检测到脸、异常边框或非有限预测时，该源帧质量置零并记录原因；只有质量为正的源帧参与聚合。采样帧的支持区间由相邻采样PTS中点划分，是分段常值近似，不意味着原视频每一帧都被检测。",10.5)
    D.table(["已具备","尚不能保证"],[
        ["原帧身份、边框、检测分数和跟踪重置可回查","最大脸/IoU规则不保证找到实际说话人"],
        ["无脸与异常数值具有显式掩码","检测分数不是情感置信度，也不是完整遮挡评分"],
        ["保存原始AU与五点地标","未提取身体姿态，未实现独立的严重遮挡判定"]],[8.5,8.5],title="视觉模块的能力边界")
    D.note("本轮复核已发现鸟羽误报：8帧具有正检测质量但并非人脸。独立排除掩码已保存，当前主特征仍含这些误报，详见第19节。")

    D.page("6  共同时间坐标与源身份","统一的是时间解释和索引关系；各模态原生序列长度可以不同。")
    D.p("对样本s定义片段内共同时间轴[0,Dₛ]。共同原点由音视频实际解码时间确定；视频位置使用原始帧PTS，音频使用首个解码音频帧偏移与重采样采样下标。文字通过强制对齐获得落在该坐标中的词区间。",10.5)
    D.equation("τₙᴬ = δᴬ + n / 16000 ,     τⱼⱽ = PTSⱼ − t₀")
    D.equation("rⱼ⁽ᵐ⁾ = ( sample_id, m, original_index, [aⱼ,bⱼ), validⱼ, qⱼ )")
    D.p("源记录r同时保存样本、模态、原始索引、时间区间、有效性和质量。文本另有原文字符跨度与token映射；视觉另有原帧号和检测边框。这样，同一输出位置可反查到明确的原始素材范围。",10.5)
    D.image(figs["timeline"],"时间组织示意图（非某条样本实测数据）；原生采样密度各异，输出锚点共用物理时间。")
    D.note("有效词顺序须保持映射单调，但无效词时间不参与单调性证据。保留原生索引，可区分真实时钟错位、采样量化及对齐器失败。")

    D.page("7  时间重叠与质量加权的数学模型","主方案采用词区间作为锚点，保留长度至少80 ms的未分配时间段。")
    D.p("设第i个目标区间为Iᵢ=[sᵢ,eᵢ)，第m个分量的原生源区间为Jⱼ=[aⱼ,bⱼ)，源特征为xⱼ，质量为qⱼ。源区间无效、特征非有限或质量不大于零时，不参与聚合。",10.5)
    D.equation("ℓᵢⱼ = max(0, min(eᵢ,bⱼ) − max(sᵢ,aⱼ))")
    D.equation("wᵢⱼ = ℓᵢⱼ qⱼ / ∑ᵣ ℓᵢᵣ qᵣ")
    D.equation("μᵢ = ∑ⱼ wᵢⱼ xⱼ ,     σᵢ = √[ ∑ⱼ wᵢⱼ (xⱼ − μᵢ)² ]")
    D.p("当分母为零时，均值和标准差写为零，同时掩码为false。零值本身可能是合法观测，因此必须依靠掩码判断缺失。情感表示使用均值；LLD与视觉同时保留区间内均值和标准差。",10.5)
    D.equation("coverageᵢ = | Iᵢ ∩ U_G | / |Iᵢ|")
    D.p("其中G为有效源索引集合，U_G为这些源区间的并集。覆盖率按时间并集计算，避免重叠声学窗口被重复计数；它与质量加权系数不是同一概念。源集合及归一化权重以CSR形式保存，可独立重算聚合结果。",10.5)
    D.table(["规则","实现含义"],[
        ["词区间","保留全部原文词；无效时间保留语义但不分配音视频"],
        ["未分配间隙","相邻有效词区间之间≥80 ms的时间保留；word_index=-1"],
        ["阈值容差","80 ms比较使用1 μs容差，避免浮点舍入造成不一致"],
        ["不擅自分类","间隙可能包含停顿、背景声、笑声或对齐失败，不自动赋予声音类别"]],[4.1,12.9],title="目标区间构造与边界处理")

    D.page("8  输出文件、掩码与证据回溯","NPZ保存数值和映射；配套JSON保存文字、帧身份、质量记录与配置引用。")
    D.table(["字段","形状/类型","用途"],[
        ["text / audio / vision","[L,768/818/56]，FP16","对齐后的三模态特征"],
        ["timestamps","[L,2]，FP64","共同时间坐标中的起止秒数"],
        ["length / valid_mask","标量 / [L]","真实长度；批量padding位置为false"],
        ["time_valid_mask","[L]","该位置时间是否可用于聚合"],
        ["modality_mask","[L,3]","文本、音频、视觉的可用状态"],
        ["component_mask","[L,4]","文本/e2v/LLD/视觉分别可用"],
        ["coverage","[L,3]","e2v/LLD/视觉源的有效时间覆盖率"],
        ["word_indices","[L]","原词索引；间隙为−1"],
        ["source_<m>_intervals","[Sₘ,2]","各源帧的原生时间窗"],
        ["map_<m>_offsets","[L+1]","CSR每行起止偏移"],
        ["map_<m>_indices / weights","[nnz]","参与聚合的源行号与权重"]],[6.1,4.3,6.6],title="主NPZ文件字段规范",size=8.7)
    D.h("从输出位置回到原素材")
    D.p("选定样本和目标位置i，读取offsets[i:i+2]，再取对应indices与weights。源行号可定位缓存特征与源时间窗；视觉源行继续映射到原始frame_index和PTS，文本通过word_indices定位原文字符跨度。",10.5)
    D.code("from q1.dataset import load_sample, collate\ns = load_sample('outputs/q1/features/<key>.npz')\nbatch = collate([s])\n# 三模态同一 L_max；padding 与观测有效性分别看掩码")
    D.note("同一样本的输出位置数一致，来自共同时间区间；并非先把原生文本、声学帧和视频帧机械拉伸成相同长度。原始源序列及其映射均保留。")

    # Later sections are appended below.
    add_results(D,samples,summary,probe,reports,vad,runconfig,figs,root,study,totalwords,validwords,positions,nativeframes,detected)
    D.fill_toc();path=OUT/(NAME+".docx");D.doc.save(path)
    sources=[root/"summary.csv",root/"manifest.json",root/"validation.json",root/"run_config.json",root/"analysis/probe_v2/report.json",root/"analysis/alignment_vad.json",study/"verification.json"]+[study/n/"report.json" for n in reports]
    meta={"docx":path.name,"figures":D.figures,"tables":D.tables,"equations":D.eq,"toc_titles":D.toc,
          "sources":{str(p.relative_to(PROJECT)):file_hash(p) for p in sources},
          "generator_sha256":file_hash(__file__),"figures_code_sha256":file_hash(figures.__file__),"equations_code_sha256":file_hash(equations.__file__),
          "human_boundary_truth":False,"known_false_face_frames_in_primary_features":8}
    (OUT/"document_manifest.json").write_text(json.dumps(meta,ensure_ascii=False,indent=2)+"\n")
    print(path);print(json.dumps({k:meta[k] for k in ["figures","tables","equations"]},ensure_ascii=False))
    return path


def add_results(D,samples,summary,probe,reports,vad,runconfig,figs,root,study,totalwords,validwords,positions,nativeframes,detected):
    v=reports["validation"];a=reports["anchors"];p=reports["sampling"];soft=reports["soft"];review=reports["review"]
    D.page("9  全量提取结果与时间覆盖","全部100条主产物通过结构与时间映射核验；形式有效性与真实准确性分别报告。")
    D.table(["统计项","当前结果","解释"],[
        ["完整样本 / 原视频组","100 / 37","全部视频与标注一一对应"],
        ["总时长",f"{v['clip_seconds']:.6f} 秒","共同时间轴长度之和"],
        ["原文词 / 时间有效词",f"{totalwords} / {validwords}",f"按词加权形式有效率{validwords/totalwords:.2%}"],
        ["输出位置 / 未分配位置",f"{positions} / {positions-totalwords}","包含词位置与≥80 ms间隙"],
        ["视觉采样帧 / 正质量帧",f"{nativeframes} / {detected}","正质量不等于人工确认人脸"],
        ["整段数字静音","2 条","音频与词时间作无效处理"],
        ["结构可疑 / 扩展复核集合","11 / 18 条","后者加入低VAD覆盖与不可评分情况"],
        ["主特征及映射体积",f"{read_json(root/'validation.json')['feature_bytes']/1e6:.3f} MB","十进制单位；未计工作缓存和本报告"]],[5.1,4.2,7.7],title="主流程全量结果概览")
    D.table(["覆盖对象","有效源时长/秒","已映射/秒","保留比例"],[
        [m,f"{r['valid_seconds']:.3f}",f"{r['mapped_seconds']:.3f}",f"{r['retained_ratio']:.4%}"] for m,r in v["source_support"].items()
    ],[4.3,4.2,4.2,4.3],title="各分量有效源支持的时间并集保留情况")
    D.p(f"全部目标区间的并集覆盖{v['mapped_seconds']:.6f}秒，占原视频总时长{v['mapped_seconds']/v['clip_seconds']:.4%}。未覆盖总计{v['unmapped_seconds']:.3f}秒，最长单段为{1000*v['longest_gap_seconds']:.1f} ms，小于80 ms间隙保留阈值。",10.5)
    D.note("时间覆盖率包含无文本的未分配区间；词时间有效率只反映边界形式检查。两者均不能替代人工标注下的词身份与边界准确率。")

    D.page("10  典型样本：同轴展示与原帧回查","样本 -3g5yACwYnA$_$13｜5.500秒｜15词｜22个输出位置。")
    s=samples[0];D.p(s["raw_text"],10.5,color=TEAL)
    D.image(figs["typical"],"词序号、波形、F0、AU原始输出与三张原帧同轴展示；红叉为无效词时间。",width=16.1)
    rows=[];z=s["export"]
    for wi in [0,5,11,13]:
        i=int(np.flatnonzero(z["word_indices"]==wi)[0]);w=s["words"][wi];counts=[]
        for m in ("emotion","acoustic","vision"):
            lo,hi=z[f"map_{m}_offsets"][i:i+2];counts.append(int(hi-lo))
        lo,hi=z["map_vision_offsets"][i:i+2];idx=z["map_vision_indices"][lo:hi]
        frames=s["vision"]["frame_indices"][idx]
        rows.append([f"{wi+1} / {w['text']}",f"{w['start']:.2f}—{w['end']:.2f}"," / ".join(map(str,counts)),f"{int(frames.min())}—{int(frames.max())}" if len(frames) else "无",str(bool(w["time_valid"]))])
    D.table(["图中词序号 / 原词","自动区间/秒","源数 e2v/LLD/V","原视频帧号范围","时间有效"],rows,[4.5,3.0,4.0,3.3,2.2],title="示例词位置的实际来源映射",size=8.4)
    D.p("词序号为展示用1起始编号；存储word_indices从0开始。表中的帧号范围只摘要显示，实际参与的稀疏帧索引与权重完整保存在CSR映射中。",9,color=GRAY)

    D.page("11  验证体系与VAD一致性检查","采用结构、时间、数值、扰动和人工证据五类互补验证。")
    D.table(["验证层次","检查内容","结果/边界"],[
        ["完整性与结构","ID、维度、长度、有限值、掩码和padding","100/100通过"],
        ["源身份与时间","原帧PTS、索引、词映射单调及长间隙","100/100通过"],
        ["聚合可重建性","独立核验源集合、重叠质量权重与FP16导出","主产物及500份替代视图通过"],
        ["实验实现","分组隔离、指标复算、检查点重放","35项测试；75模型重放通过"],
        ["学习映射约束","留出注意力的有效性与时间窗外零权重","25份映射核验通过"],
        ["人工真实性","听音词边界、真实人脸/说话人确认","词边界尚未人工标注"]],[3.5,8.7,4.8],title="验证证据及其适用范围",size=8.8)
    D.p("Silero VAD独立于转写读取音频，将语音活动与有效词区间投影到10 ms网格。Precision衡量词时间中的语音比例；Recall衡量语音中被词时间覆盖的比例；IoU衡量两者交并比。它能发现明显不一致，不能确认每个词的身份。[D5]",10)
    D.table(["可评分样本","Precision中位数","Recall中位数","IoU中位数","待复核"],[[str(vad["scored"]),f4(vad["all"]["precision_median"]),f4(vad["all"]["recall_median"]),f4(vad["all"]["iou_median"]),"18条"]],[3.0,3.8,3.6,3.6,3.0],title="VAD辅助核查结果")
    D.image(figs["vad"],"VAD覆盖分布；4条不可评分样本未当作零准确率加入统计。")
    D.note("VAD使用32 ms输入块和0.5阈值；循环平移检验保留词区间布局。语音几乎覆盖整段时该对照区分力有限，p值不用于删除样本。")

    D.page("12  辅助实验：统一协议与公平比较","情感预测探针用于检查表示是否含有可用信息，不代替第一问的时间真实性验证。")
    D.table(["环节","设置与目的"],[
        ["输入","附件1的100条、37个原视频及已有标签；不引入其他情感数据集"],
        ["外层验证","按video_id留一，共37折，防止同原视频片段跨训练/测试"],
        ["内层选择","外层训练集内5折GroupKFold；17档alpha从0.01到10⁶"],
        ["回归 / 三分类","Ridge按内层MAE选alpha；独热Ridge按内层Macro-F1独立选择"],
        ["预处理","每个训练折独立拟合缺失均值填补与标准化"],
        ["特征权重","按完整家族参考维度缩放：文本/e2v=768，LLD=50，视觉=56"],
        ["标签置换","499次视频组内置换，每次重做内层选择；每指标17项Holm校正"],
        ["区间估计","2000次按原视频的成对Bootstrap；区间未作多重比较校正"]],[3.4,13.6],title="17组维度与融合探针的实验协议")
    D.p("分组验证遵循同一视频的片段共同留出的原则。[R6] 标准化、填补和正则参数选择均在训练侧进行。缺失分量记为NaN后在训练折填补，避免用测试统计量补全。各特征家族使用固定缩放分母，消融删除坐标时不改变保留坐标的权重。",10.5)
    D.h("三个实验族不能混写")
    D.table(["实验族","片段读出与划分","可回答的问题"],[
        ["维度/融合探针","有效词位置均值；37个外层折","维度及模态组合的探索性作用"],
        ["锚点对照","包含间隙；等权/时长加权；37折","不同区间统计尺度的差异"],
        ["软对齐对照","原生编码；5折×3种子；固定训练轮数","同一小模型内注意力约束的作用"]],[3.4,8.2,5.4],title="实验口径的区别")
    D.note("组内置换保留每个原视频的标签多重集；仅22组连续标签可在组内变化。它检验特定的组内可交换假设，不能解释为对所有视频间标签关系的全面检验。")

    D.page("13  维度消融与融合结果","所有数值为完整100条的外层留出预测；Acc与F1均为含中性类的三分类指标。")
    names={"text_modernbert":"文本 ModernBERT","text_bert":"文本 BERT对照","audio_768_aligned":"音频768（词位置）","audio_793_aligned":"音频793（词位置）","audio_818_aligned":"音频818（词位置）","audio_768_clip":"音频768（整段）","audio_818_clip":"音频818（整段）","egemaps_50_aligned":"LLD均值＋标准差","vision_28_aligned":"视觉28（词位置）","vision_56_aligned":"视觉56（词位置）","vision_56_clip":"视觉56（整段）","audio_visual_aligned":"音频＋视觉","text_audio_aligned":"文本＋音频","text_visual_aligned":"文本＋视觉","fusion_aligned":"完整三模态（词位置）","fusion_no_std":"三模态去除标准差","fusion_clip":"完整三模态（整段）"}
    rows=[]
    for n,r in probe["results"].items():rows.append([names[n],r["dims"],f4(r["mae"]),f4(r["pearson"]),f4(r["accuracy3"]),f4(r["macro_f1_3"]),f"{r['p_macro_f1_3_holm']:.3f}"])
    D.table(["配置","维度","MAE↓","Pearson↑","Acc3↑","Macro-F1↑","F1校正p"],rows,[5.0,1.25,2.0,2.0,2.0,2.35,2.4],title="修正版探针的17组完整结果",size=8.0)
    D.image(figs["probe"],"部分维度与融合配置的Macro-F1；柱高差异不自动等于显著增益。")
    D.p("音频768→818的Macro-F1提升约0.1195，探索性95%差值区间[0.0247,0.2092]；但818维分类的Holm校正p=0.068。793→818及视觉28→56的差值区间均跨0。全部17组的各指标组内置换检验经校正均未达到0.05。",9.4)
    D.note("768/818/56是定义清楚、可复现的设计选择，当前证据不支持称其最优。词位置融合相对整段融合也未建立稳定优势。")

    D.page("14  锚定区间对照：词、秒与固定段数","全量100条；文本全片内容固定，只改变音视频聚合及其统计尺度。")
    D.p("比较词区间＋间隙、0.25/0.5/1秒固定区间和全片均匀50段。每种区间分别采用锚点等权均值与按区间时长加权的片段读出，共10组配置。采用37折外层、5折内层、499次组内置换和2000次视频Bootstrap。",10.5)
    modes=a["protocol"]["modes"];labels=["词＋间隙","固定0.25秒","固定0.5秒","固定1秒","均匀50段"]
    rows=[]
    for mode,label in zip(modes,labels):
        u,d=(a["results"][mode+"__"+pool] for pool in ("uniform","duration"))
        rows.append([label,f4(u["mae"]),f4(u["macro_f1_3"]),f4(d["mae"]),f4(d["macro_f1_3"])])
    D.table(["锚点","等权MAE","等权F1","时长加权MAE","时长加权F1"],rows,[4,3.1,3.1,3.4,3.4],title="五种锚点与两种读出的结果")
    D.image(figs["anchors"],"锚点比较；固定50段等长，因此两种读出结果相同。")
    D.p("固定50段相对词区间方案，MAE和Macro-F1成对差值区间均跨0；目前没有替换依据。固定0.25秒时长加权的F1相对词区间略有不利信号，但属于未校正的探索性比较。全部10配置的置换检验经Holm校正均未通过0.05阈值。",10.5)
    D.note("均值读出会丢失序列顺序，不能据此否定固定锚点在其他时序模型中的作用。固定网格也不能修复无效词边界。500份替代特征与映射均已独立重建核验。")

    D.page("15  采样率扰动：坐标准确与表示稳定性","8条预先选择的诊断样本，5/8/10/12/25 Hz；使用固定的训练折读出器。")
    D.p("样本选择依据为首条参考、两条静音、无脸、短/长片段、低VAD覆盖与视觉变化，不使用情感标签。实验按固定秒网格选择最近的原始PTS帧并去重，重新检测/提取视觉特征；以新网格10 Hz为基准，另保留旧10 Hz缓存作为对照。",10.5)
    rows=[]
    for rate in [5,8,10,12,25]:
        z=p["summary"][str(rate)];rows.append([rate,f"{z['effective_hz']:.3f}",f"{z['relative_feature_change']:.3f}",f"{1000*z['mean_abs_centroid_shift_seconds']:.2f}",f"{z['mean_abs_prediction_change']:.4f}",f"{z['class_flips']}/8"])
    D.table(["目标Hz","实际Hz中位数","相对特征变化","时间偏移/ms","平均|回归变化|","类别变化"],rows,[2,3.3,3.2,2.7,3.2,2.6],title="相对网格10 Hz的扰动结果",size=8.6)
    D.image(figs["sampling"],"时间偏移与特征变化；仅在双方均有有效观测的位置比较。")
    D.p("所有40份新采样缓存的原帧PTS身份误差均为0。12 Hz与25 Hz各有1条分类变化，说明时间身份稳定，并不保证特征或预测逐值不变。不同取帧密度也会改变跟踪历史和分段常值支持近似。",10.5)
    D.note("旧最小间隔规则受帧量化影响：约24 fps视频的名义10 Hz采样可接近8 Hz。新网格采样仅用于实验；目标频率高于原帧率时不生成插值帧。这8条诊断结果不能外推为100条全量稳健性保证。")

    D.page("16  原生编码与局部软对齐的实现","作为独立小模型对照；原生特征、质量掩码与物理CSR映射保持各自身份。")
    D.p("文本768维、emotion2vec768维、LLD25维与视觉28维分别经Linear→LayerNorm→GELU映射到16维潜在空间。每0.5秒设置查询锚点，用时间重叠的文本潜变量形成查询；完整文本语义分支仍包含时间无效词。",10.5)
    D.equation("hⱼ⁽ᵐ⁾ = GELU( LayerNorm( Wₘ xⱼ⁽ᵐ⁾ + bₘ ) )")
    D.equation("scoreᵢⱼ = QᵢKⱼᵀ / √16 − 0.5(Δtᵢⱼ / 0.25)² + log(qⱼ)")
    D.equation("αᵢⱼ = softmaxⱼ(scoreᵢⱼ),     |Δtᵢⱼ| ≤ 0.5 秒")
    D.p("局部窗外、质量为零和padding源禁止参与注意力。若一个查询没有有效源，则返回全零权重和无效标记；不让softmax在全缺失条件下虚构观测。各分量上下文的均值/标准差、文本分支及可用性标记共同进入小型读出头。",10.5)
    D.table(["对照","相对带惩罚局部注意力的改变"],[
        ["hard_overlap：硬重叠","只用时间重叠×质量，忽略内容Q/K分数"],
        ["time_kernel：固定时间核","只用局部高斯时间权重和质量，忽略内容分数"],
        ["local_attention：局部注意力","内容分数＋±0.5秒窗＋时间距离惩罚"],
        ["local_no_penalty：局部无惩罚","保留局部窗口，移除时间距离惩罚"],
        ["global_attention：全局注意力","保留有效性掩码，移除局部窗口与距离惩罚"]],[6.0,11.0],title="五组匹配对照")
    D.p("训练协议预先固定：5折GroupKFold×3个种子（2026/2027/2028）×5方法，共75次；80轮全训练集批次，AdamW学习率0.001、权重衰减0.01、dropout=0.1。损失为训练折标准化标签的MSE＋训练类别加权交叉熵；不根据测试表现早停或调参。",10)

    D.page("17  软对齐结果：收益与不确定性","同一5折协议内比较；误差条为3个种子的标准差，不是95%置信区间。")
    labels={"hard_overlap":"硬重叠","time_kernel":"固定时间核","local_attention":"局部注意力","local_no_penalty":"局部无惩罚","global_attention":"全局注意力"}
    rows=[]
    for n,z in soft["results"].items():rows.append([labels[n],f"{z['mae']['mean']:.4f}±{z['mae']['seed_std']:.4f}",f4(z["pearson"]["mean"]),f4(z["accuracy3"]["mean"]),f"{z['macro_f1_3']['mean']:.4f}±{z['macro_f1_3']['seed_std']:.4f}"])
    D.table(["方法","MAE均值±SD","Pearson均值","Acc3均值","Macro-F1均值±SD"],rows,[3.3,4.1,3.0,2.5,4.1],title="五种软/硬对齐对照结果",size=8.8)
    D.image(figs["soft_results"],"神经对照结果；三种子统一汇总，不选择表现最好的种子。")
    comparisons=[]
    for n in ("hard_overlap","time_kernel","local_no_penalty","global_attention"):
        c=soft["paired_bootstrap"]["local_attention - "+n]
        comparisons.append(["局部−"+labels[n],f"{c['mae']['difference']:+.4f}",f"[{c['mae']['ci95'][0]:+.4f}, {c['mae']['ci95'][1]:+.4f}]",f"{c['macro_f1_3']['difference']:+.4f}",f"[{c['macro_f1_3']['ci95'][0]:+.4f}, {c['macro_f1_3']['ci95'][1]:+.4f}]"])
    D.table(["比较","MAE差","MAE差95%区间","F1差","F1差95%区间"],comparisons,[4.1,2.1,4.35,2.1,4.35],title="按原视频成对Bootstrap的探索性差值",size=8.0)
    D.note("局部注意力没有明确超过硬重叠；相对全局注意力有探索性优势，但区间未作多重比较校正。同折训练均值回归基线MAE为0.5968，小模型回归没有显示明确优势。")

    D.page("18  注意力映射的追溯与解释边界","展示预先确定的留出样本：第0折、种子2026；不按预测好坏挑选。")
    D.image(figs["attention"],"同一留出样本的视觉源映射；横纵轴均为物理时间，颜色为权重。")
    D.table(["记录","可核查内容"],[
        ["查询锚点","anchor_times：0.5秒区间及其中心位置"],
        ["原生源身份","source_indices、source_intervals；视觉另含原帧号与PTS"],
        ["模型身份","训练/留出索引、fold、seed、mode及文件哈希"],
        ["权重约束","权重非负；有观测时归一化；无效源和局部窗外为零"],
        ["独立重放","75个保存模型可重放相同预测；25份注意力映射通过约束检查"]],[4.5,12.5],title="学习映射的审计字段")
    D.p("硬重叠对应一个直接可解释的物理区间加权过程；局部注意力在邻近时间范围内调整关联强度。两者共享时间坐标，但学习权重并不表示原词时间被重新标注，也不意味着较大的注意力权重必然对应真实情绪因果证据。",10.5)
    D.p("当前图示只说明模型遵守局部时间约束。问题3要求的模态作用程度和关键证据解释，仍需在其独立任务与数据协议中验证，不能用此注意力热图代替完整的问题3方法。",10.5)
    D.note("这组实验的作用是检验“允许邻近异步关联是否有帮助”。当前结果支持保留局部约束作为候选设计，但不支持将其作为第一问已获证实的性能创新。")
    add_quality_and_appendices(D,samples,summary,probe,reports,vad,runconfig,figs,root,study)


def add_quality_and_appendices(D,samples,summary,probe,reports,vad,runconfig,figs,root,study):
    review=reports["review"]
    D.page("19  已知质量问题：鸟羽人脸误报","结构核验正确仍可能伴随观测语义错误；本例需要单独的质量复核证据。")
    D.p("样本 -NFrJFQijFE$_$2 为鸟类画面。原视觉模型在57个采样帧中的8帧返回正质量人脸观测。进一步查看全部8个检测框发现，它们都位于羽毛区域，不是真实人脸。下图选取其中4个框展示。",10.5)
    D.image(figs["false_face"],"鸟羽误报的实际原帧与原检测框（8例中展示4例）；不是模拟图。")
    D.table(["项目","记录与当前处理"],[
        ["误报原帧号","45、48、60、63、66、69、96、117"],
        ["证据记录","scene_review.json：原视频哈希、复核类型与排除原帧号"],
        ["独立掩码","review_excluded_mask；usable_after_review_mask"],
        ["主产物状态","8帧误报仍在原主特征中；本轮实验未应用事后掩码"],
        ["后续修正","在源帧层屏蔽、重新聚合视觉均值/标准差/CSR并再核验"]],[4.0,13.0],title="误报的证据、修正材料与未完成部分")
    D.note("复核排除掩码只纠正已经查明的8帧，不能证明其余帧都正确。主体脸跟踪、遮挡程度、低光与真实说话人身份仍是当前视觉处理的局限。")

    D.page("20  18条可疑样本与人工复核材料","已提供播放与标注页面；263个词的人工边界字段保持空白。")
    cases=read_json(study/"review/cases.json");lookup={s["id"]:s for s in samples}
    cats={"digital_silence_confirmed_by_waveform":"整段数字静音","no_speech_detected_by_vad_requires_listening":"VAD未检出语音","low_speech_time_coverage_requires_word_boundary_review":"低语音覆盖","structurally_invalid_words_require_review":"词时间结构异常"}
    rows=[]
    for c in cases:
        s=lookup[c["sample_id"]];valid=sum(w["time_valid"] for w in s["words"])
        reason=cats[c["automatic_category"]]+("；鸟羽误报" if c["review_excluded_frames"] else "")
        rows.append([c["sample_id"],f"{c['duration']:.3f}",f"{valid}/{len(s['words'])}",f4(c["vad_recall"]),reason])
    D.table(["样本ID","时长/秒","有效词/总词","VAD Recall","复核原因"],rows,[5.4,2.0,2.8,2.3,4.5],title="待复核集合的全量清单",size=8.2)
    D.p("复核页面包含原转写、缓存音频、原视频、波形/VAD/词区间、每条3张原帧及逐词标注导出。AI已查看54张抽查帧，并额外核查鸟类片段全部8个有效检测框；未逐帧确认所有片段，也未进行人工听音词边界标注。",10)
    D.p("人工标注使用“缓存音频播放秒数＋audio_offset”作为共同时间；原视频播放仅辅助查看。标注者填写匿名代号、词起点和终点，导出CSV后运行评分。未知/重复词、半个边界、越界和缺失复核者均会被拒绝。",10)
    D.note("人工边界MAE、P90及100 ms命中率目前均不可报告。两条数字静音由波形确认，但不能凭已有转写推测静音中的词边界。")

    D.page("21  运行流程、缓存与交付文件","所有命令在项目根目录执行；主流程与分析产物分别管理。")
    D.h("主流程")
    D.code(".venv-q1/bin/python -m q1 doctor\n.venv-q1/bin/python -m q1 manifest\n.venv-q1/bin/python -m q1 run\n.venv-q1/bin/python -m q1 validate\n.venv-q1/bin/python -m q1 preview --id=-3g5yACwYnA__13")
    D.h("辅助验证与报告")
    D.code(".venv-q1/bin/python -m q1.analysis.alignment_quality\n.venv-q1/bin/python -m q1.analysis.probe --threads 2\n# 后续模块位于 q1.analysis.time_study：\n# validate / review / sampling / anchors / soft\n.venv-q1/bin/python -m q1.analysis.time_study.verify_experiments\n.venv-q1/bin/python -m q1.analysis.time_study.report\n.venv-q1/bin/python -m unittest discover -s q1/tests -v")
    D.p("缓存签名包括输入视频与原文、配置、主代码、依赖版本、上游产物及输出校验和。异常阶段不会以随机向量或全零结果冒充成功；损坏或过期缓存不复用。全量流程失败会返回非零状态，并保留失败样本记录。",10.5)
    D.table(["路径（相对outputs/q1）","用途"],[
        ["features/","正式100条特征与来源映射"],["manifest.json / summary.csv","输入身份与全量汇总"],
        ["run_config.json / events.jsonl","实际配置、环境及处理日志"],["cache/","工作用波形、原生特征、时间和阶段签名"],
        ["analysis/probe_v2/","修正版维度与融合探针结果"],["analysis/time_study/","时间核验、扰动、锚点、软对齐与复核"],
        ["documentation/","本Word、PDF预览、图件和构建来源清单"]],[8.3,8.7],title="主要产物位置")
    D.note("当前主特征与映射约9.675 MB；工作缓存和全部实验检查点不应整体作为竞赛附件。题目总附件≤50 MB还要容纳问题2/3材料。本报告使用相对路径与匿名元数据。")

    D.page("22  结论、限制与后续优先级","当前证据支持保留可解释的词区间＋间隙主方案。")
    D.h("已经成立的结论")
    D.p("已构建覆盖100条视频的三模态特征提取流程，明确768/818/56维表示的来源与统计定义；建立共同时间坐标、原生索引、有效掩码及可重建的聚合映射。主产物结构核验、时间身份检查及数值重建均通过。",10.5)
    D.p("实验已覆盖维度与融合消融、五类锚点、采样扰动、五种软/硬对齐及留出映射审计。固定50段没有明确超过词区间；局部注意力没有明确超过硬重叠。保持主方法简洁有当前实验依据。",10.5)
    D.h("不能据此宣称的结论")
    D.table(["不宜宣称","原因"],[
        ["对齐准确率99.9706%","该数值是含间隙的时间并集覆盖率"],
        ["768/818/56是最优维度","分类有探索性信号，但多项校正检验和成对比较证据不足"],
        ["局部软对齐显著提升主方法","相对硬重叠的MAE/F1差值区间均跨0"],
        ["所有有效视觉帧都是真实人脸","已发现8帧羽毛误报；检测分数不能替代语义复核"],
        ["注意力已实现问题3可解释性","此处只验证局部关联约束，未完成专项解释任务"]],[6.5,10.5],title="报告结论的明确边界")
    D.h("优先完成的收尾工作")
    D.p("第一，应用已知视觉误报的源帧排除掩码，重新聚合主视觉特征与映射，并重新验证。第二，从正常与异常片段中取得人工听音词边界，报告真实边界误差，避免只标注容易样本。第三，将完整特征、图表、全量汇总、模型版本与运行说明整理为可提交材料。",10.5)
    D.note("本报告描述截至2026年9月24日的当前状态。自动实验已完成；人工词边界及已知误报合入主产物尚未完成。后续修改输入或掩码后，相关统计与实验应重新运行并形成新版本报告。")

    # Appendix A: all 100 manifest-order samples, 25 per page.
    for page in range(4):
        title="附录A  全部100条样本汇总" if page==0 else f"附录A（续{page}）  样本{page*25+1}—{(page+1)*25}"
        D.page(title,"按原始标注清单顺序；每行同时对应文本、语音、视觉三种模态。",toc=page==0)
        D.p("统一维度为768/818/56；对齐粒度为词区间＋≥80 ms未分配间隙。D是原始片段共同时间轴时长；L是最终序列长度；有效比例T/A/V以全部L位置为分母，包含无文本间隙。异常标记“疑”指主流程结构可疑，“静”指数字静音。",9)
        rows=[]
        for i,(_,r) in enumerate(summary.iloc[page*25:(page+1)*25].iterrows(),start=page*25+1):
            flags="、".join((["静"] if r.silent_audio else [])+(["疑"] if r.alignment_suspect else [])) or "—"
            rows.append([i,r.sample_id,f"{r.duration:.3f}",int(r.words),int(r.positions),"768/818/56",f"{100*r.valid_text_ratio:.0f}/{100*r.valid_audio_ratio:.0f}/{100*r.valid_vision_ratio:.0f}",flags])
        D.table(["序","样本ID","D/秒","词数","L","维度T/A/V","有效T/A/V%","标记"],rows,[.9,4.8,1.55,1.0,.9,2.9,3.55,1.4],title=f"全部样本清单（{page*25+1}—{(page+1)*25}）",size=8.0)
        D.p("样本ID保留原始“video_id$_$clip_id”格式；文件名将分隔符转换为“__”。完整数值、源时间窗、质量及映射见对应NPZ/JSON。有效比例经四舍五入展示，不作为人工准确率。",8.7,color=GRAY)

    D.page("附录B  25项原生声学描述量","列名以实际openSMILE输出为准；每项在目标区间生成均值与标准差各一维。")
    names=read_json(samples[0]["directory"]/"audio_features.json")["acoustic_names"]
    descriptions=["响度","高/低频谱能量比","Hammarberg谱形指标","0—500 Hz谱斜率","500—1500 Hz谱斜率","谱通量","第1阶MFCC","第2阶MFCC","第3阶MFCC","第4阶MFCC","相对27.5 Hz参考的F0半音值","局部基频抖动","局部振幅扰动（dB）","谐噪比相关指标","H1/H2相对谱幅指标","H1/A3相对谱幅指标","第1共振峰频率","第1共振峰带宽","第1共振峰相对F0幅度","第2共振峰频率","第2共振峰带宽","第2共振峰相对F0幅度","第3共振峰频率","第3共振峰带宽","第3共振峰相对F0幅度"]
    D.table(["序号","实际字段名","便于阅读的含义"],[[i+1,n,d] for i,(n,d) in enumerate(zip(names,descriptions))],[1.2,10.3,5.5],title="eGeMAPSv02 LLD实际维度顺序",size=8.7)
    D.note("字段名中的sma3/sma3nz等后缀来自工具定义；本报告不另行改名或重新标定。LLD总计25项，不是整句级88维Functionals。")

    D.page("附录C  模型版本与关键运行参数","固定权重修订号来自q1/config.yaml；依赖版本来自实际run_config.json。")
    models=runconfig["config"]["models"];rows=[]
    for name,zh in [("align","强制对齐"),("text","文本"),("emotion","语音"),("face","视觉")]:
        m=models[name];rows.append([zh,m["id"]+"\nrevision: "+m["revision"]])
    rows.append(["BERT对照",probe["protocol"]["bert"]["id"]+"\nrevision: "+probe["protocol"]["bert"]["revision"]])
    D.table(["用途","模型与完整修订号"],rows,[2.8,14.2],title="固定模型身份",size=8.5)
    env=runconfig["environment"];items=[(k,str(env[k])) for k in ["torch","torchvision","transformers","numpy","pandas","opensmile","funasr","openface-test","timm","huggingface_hub"]]
    D.table(["依赖","实际版本","依赖","实际版本"],[[*items[i],*items[i+5]] for i in range(5)],[4.5,4,4.5,4],title="主要依赖版本",size=8.7)
    D.p("已验证Python 3.12、V100 32 GB。Qwen3使用FP16，ModernBERT使用FP32；采用eager attention以适配现有GPU。输出特征FP16、时间戳FP64；音频采样率16 kHz；视觉阈值0.8、IoU阈值0.2、长边960像素。",10)
    D.p("OpenFace包按已有安装说明适配；上游旧依赖声明与当前环境存在已知差异，复现应以q1/environment-tested.txt、install.sh及doctor检查为依据。新环境不能仅凭模型名称假定运行行为一致。",9.5,color=GRAY)

    D.page("附录D  来源、参考资料与文档构建","本地结果是数值依据；官方资料用于核对工具用途。外部资料查阅日期：2026-09-24。")
    D.table(["编号","本地材料（项目相对路径）"],[
        ["D1","E题/复杂场景下多模态情感识别的数学建模与算法设计.docx"],
        ["D2","q1/README.md、q1/config.yaml、q1/temporal.py、q1/pipeline.py"],
        ["D3","outputs/q1/manifest.json、summary.csv、validation.json、run_config.json"],
        ["D4","outputs/q1/analysis/probe_v2/report.json、conclusions.md"],
        ["D5","outputs/q1/analysis/alignment_vad.json、alignment_vad.csv"],
        ["D6","outputs/q1/analysis/time_study/各子目录report.json、verification.json"],
        ["D7","q1/analysis/time_study/scene_review.json及review/quality/排除掩码"]],[1.3,15.7],title="本报告本地证据索引",size=8.8)
    for label,url in [
        ("[R1] Qwen3-ForcedAligner官方模型卡","https://huggingface.co/Qwen/Qwen3-ForcedAligner-0.6B-hf"),
        ("[R2] ModernBERT-base官方模型卡","https://huggingface.co/answerdotai/ModernBERT-base"),
        ("[R3] emotion2vec官方代码与特征提取说明","https://github.com/ddlBoJack/emotion2vec"),
        ("[R4] openSMILE Python官方文档","https://audeering.github.io/opensmile-python/"),
        ("[R5] OpenFace 3.0官方代码与模块说明","https://github.com/CMU-MultiComp-Lab/OpenFace-3.0"),
        ("[R6] scikit-learn交叉验证与分组验证文档","https://scikit-learn.org/stable/modules/cross_validation.html")]:D.link(label,url)
    D.h("如何重建这份报告")
    D.code(".venv-q1/bin/python -m q1.analysis.time_study.documentation.build_docx")
    D.p("生成器从当前已核验JSON/CSV、原生缓存及配置读取结果；图片由本地数据绘制，表格和公式为Word可编辑内容。document_manifest.json保存来源文件SHA256、构建代码哈希及图表计数。PDF供核对排版，DOCX用于后续修订。",9.7)
    D.note("更新模型、时间映射、观测掩码或实验协议后，应先重新运行受影响的验证与实验，再重建文档；不手工更改表格数值来替代新的实验记录。")


if __name__=="__main__":run()
