"""Build the current Q1 method report from verified outputs."""
from __future__ import annotations

import shutil
import subprocess
import tempfile
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .common import atomic_json, file_hash, read_json
from .pipeline import validate
from .preview import render
from .validation import REASON_LABELS


def figures(out):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    from matplotlib.patches import FancyBboxPatch
    cjk = Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc")
    if cjk.exists():
        font_manager.fontManager.addfont(str(cjk))
        plt.rcParams["font.family"] = font_manager.FontProperties(fname=str(cjk)).get_name()
    plt.rcParams["axes.unicode_minus"] = False
    out.mkdir(parents=True, exist_ok=True)
    fig, ax = plt.subplots(figsize=(10, 7))
    ax.set(xlim=(0, 10), ylim=(0, 7)); ax.axis("off")
    def box(x, y, w, h, title, detail):
        ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=.05", facecolor="#edf5f6", edgecolor="#147d86"))
        ax.text(x+w/2, y+h*.7, title, ha="center", va="center", color="#125865", fontsize=13, weight="bold")
        ax.text(x+w/2, y+h*.3, detail, ha="center", va="center", color="#173644", fontsize=10)
    def arrow(a, b):
        ax.annotate("", xy=b, xytext=a, arrowprops={"arrowstyle":"->", "color":"#526a76", "lw":1.4})
    box(.7, 6, 8.6, .75, "附件1：100条视频＋给定英文转写", "保留样本ID、原文、原始索引及视频SHA256")
    box(.7, 4.9, 8.6, .75, "共同时间坐标", "视频使用实际PTS；音频使用采样下标＋原始偏移")
    arrow((5, 6), (5, 5.7))
    for x, title, detail in [(.2, "文本", "Qwen3补词时间\nModernBERT：768维"), (3.6, "语音", "openSMILE：25项LLD\n保留原生时间窗"), (7., "视觉", "OpenFace 3.0：28维\n原帧号、PTS、检测质量")]:
        box(x, 3.1, 2.8, 1.25, title, detail)
        arrow((x+1.4, 4.9), (x+1.4, 4.4))
        arrow((x+1.4, 3.1), (x+1.4, 2.65))
    box(.4, 1.7, 9.2, .9, "SAQW-TA：词锚点＋gap＋质量加权区间映射", "文本768维｜声学均值25＋标准差25｜视觉均值28＋标准差28")
    box(.7, .3, 8.6, .9, "正式输出：768／50／56", "有效掩码、时间覆盖、CSR来源映射、独立数值重建")
    arrow((5, 1.7), (5, 1.25))
    fig.savefig(out / "architecture.png", dpi=190, bbox_inches="tight"); plt.close(fig)
    fig, ax = plt.subplots(figsize=(10, 3.2)); ax.axis("off")
    formulas = [r"$\ell_{ij}=\max\{0,\min(e_i,b_j)-\max(s_i,a_j)\}$",
                r"$w_{ij}=\frac{\ell_{ij}q_j}{\sum_r\ell_{ir}q_r},\qquad \mu_i=\sum_j w_{ij}x_j$",
                r"$\sigma_i=\sqrt{\sum_j w_{ij}(x_j-\mu_i)^2}$"]
    for y, formula in zip([.85, .5, .15], formulas):
        ax.text(.5, y, formula, ha="center", va="center", fontsize=19, color="#173644")
    fig.savefig(out / "aggregation.png", dpi=200, bbox_inches="tight"); plt.close(fig)
    fig, ax = plt.subplots(figsize=(10, 5.2)); ax.axis("off")
    formulas = [
        r"$\mathcal{I}_s=\bigcup_{i:\,t_{si}=1}I_{si},\qquad \mathcal{J}_s^{(m)}=\bigcup_{j\in\mathcal{V}_s^{(m)}}J_{sj}^{(m)}$",
        r"$C_{si}^{(m)}=\frac{|I_{si}\cap\mathcal{J}_s^{(m)}|}{|I_{si}|}\in[0,1]$",
        r"$R_{\rm timeline}=\frac{\sum_s|\mathcal{I}_s|}{\sum_s D_s},\qquad C_m=\frac{\sum_s|\mathcal{I}_s\cap\mathcal{J}_s^{(m)}|}{\sum_s|\mathcal{I}_s|}$",
        r"$C_{A\cup V}=\frac{\sum_s|\mathcal{I}_s\cap(\mathcal{J}_s^{(A)}\cup\mathcal{J}_s^{(V)})|}{\sum_s|\mathcal{I}_s|}$",
        r"$C_{A\cap V}=\frac{\sum_s|\mathcal{I}_s\cap\mathcal{J}_s^{(A)}\cap\mathcal{J}_s^{(V)}|}{\sum_s|\mathcal{I}_s|}$"]
    for y, formula in zip([.92, .72, .51, .30, .09], formulas):
        ax.text(.5, y, formula, ha="center", va="center", fontsize=17, color="#173644")
    fig.savefig(out / "coverage.png", dpi=210, bbox_inches="tight"); plt.close(fig)
    panels = {
        "representation": [r"$X_i=(T_i,A_i,V_i,I_i,t_i,M_i,C_i,P_i)$",
                           r"$T_i\in\mathbb{R}^{768},\quad A_i=[\mu_i^{(A)},\sigma_i^{(A)}]\in\mathbb{R}^{50},\quad V_i=[\mu_i^{(V)},\sigma_i^{(V)}]\in\mathbb{R}^{56}$"],
        "quality": [r"$q_k^{(V)}=c_k\,d_k\,b_k\,f_k\,(1-r_k)$",
                    r"$q_k^{(A)}=u_k\,f_k^{(A)}\in\{0,1\}$"]}

    for name, equations in panels.items():
        fig, ax = plt.subplots(figsize=(10, 2.1)); ax.axis("off")
        for y, equation in zip([.73, .20], equations):
            ax.text(.5, y, equation, ha="center", va="center", fontsize=16, color="#173644")
        fig.savefig(out / (name + ".png"), dpi=210, bbox_inches="tight"); plt.close(fig)


class Report:
    def __init__(self):
        self.doc = Document(); section = self.doc.sections[0]
        section.page_width = Cm(21); section.page_height = Cm(29.7)
        section.top_margin = section.bottom_margin = Cm(1.8)
        section.left_margin = section.right_margin = Cm(2)
        normal = self.doc.styles["Normal"]
        normal.font.name = "Noto Serif CJK SC"; normal.font.size = Pt(10.5)
        normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Noto Serif CJK SC")
        normal.paragraph_format.space_after = Pt(7); normal.paragraph_format.line_spacing = 1.16
        for level, size in [(1, 19), (2, 13)]:
            style = self.doc.styles[f"Heading {level}"]
            style.font.name = "Noto Sans CJK SC"; style.font.size = Pt(size)
            style.font.color.rgb = RGBColor.from_string("147D86")
            style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Noto Sans CJK SC")
        props = self.doc.core_properties
        props.author = ""; props.last_modified_by = ""
        props.title = "第一问：多模态特征提取与时序对齐"
        footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run("第一问 · 768／50／56    ").font.size = Pt(8)
        page = OxmlElement("w:fldSimple"); page.set(qn("w:instr"), "PAGE"); footer._p.append(page)
        self.tables = self.images = 0

    def p(self, text):
        return self.doc.add_paragraph(text)

    def page(self, title):
        self.doc.add_page_break(); self.doc.add_heading(title, 1)

    def table(self, headers, rows, widths=None, size=9):
        self.tables += 1
        t = self.doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        widths = widths or [17/len(headers)]*len(headers)
        for col, width in zip(t.columns, widths): col.width = Cm(width)
        for i, values in enumerate([headers]+list(rows)):
            row = t.rows[0] if i == 0 else t.add_row()
            trpr = row._tr.get_or_add_trPr(); trpr.append(OxmlElement("w:cantSplit"))
            if i == 0: trpr.append(OxmlElement("w:tblHeader"))
            for j, value in enumerate(values):
                cell = row.cells[j]; cell.width = Cm(widths[j]); p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(3); p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.line_spacing = 1.08
                r = p.add_run(str(value)); r.font.size = Pt(size); r.bold = i == 0
                r.font.color.rgb = RGBColor.from_string("FFFFFF" if i == 0 else "173644")
                sh = OxmlElement("w:shd"); sh.set(qn("w:fill"), "147D86" if i == 0 else ("EFF5F6" if i % 2 else "FFFFFF"))
                cell._tc.get_or_add_tcPr().append(sh)
        self.p("")

    def image(self, path, caption, width=16.7):
        self.images += 1
        p = self.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Cm(width))
        self.p(f"图{self.images}  {caption}")


def build_report(cfg, samples):
    root = Path(cfg["output_dir"]); out = root / "report"; out.mkdir(exist_ok=True)
    result = validate(cfg, samples)
    if result["complete"] != len(samples): raise ValueError("须先完成全部样本验证")
    summary = pd.read_csv(root / "summary.csv")
    figures(out / "figures")
    example = next(s for s in samples if s["key"] == "-3g5yACwYnA__13")
    typical = render(cfg, example)
    example_dir = root / "cache" / example["key"]
    names = read_json(example_dir / "audio.json")["acoustic_names"]
    record = read_json(root / "run_config.json")
    coverage, alignment, availability = (result[k] for k in ("temporal_coverage", "word_alignment", "modality_availability"))
    D = Report(); D.doc.add_heading("第一问：多模态特征提取\n与时序对齐", 0)
    D.p("SAQW-TA：语义锚点质量加权时序对齐方法")
    D.p("Qwen3 ForcedAligner + ModernBERT + openSMILE + OpenFace 3.0")
    D.p("结果快照：" + datetime.now(timezone.utc).strftime("%Y-%m-%d") + "；全部数值来自当前代码与正式产物。")
    D.table(["完整样本", "三模态维度", "目标时间轴保留率"], [[f"{result['complete']}/{result['expected']}", "768／50／56", f"{coverage['timeline_retention']['ratio']:.4%}"]])
    D.image(out / "figures/architecture.png", "主流程与输出接口。", 15.5)
    D.p(f"任务目标是定义可复现、可回溯的三模态时序特征。全部100条样本保留。{coverage['timeline_retention']['ratio']:.4%}衡量有效词区间与保留gap构成的目标时间轴完整性，未检查源模态是否有效；源支持覆盖与词边界准确率是不同指标。")

    D.page("1  输入、工具分工与特征定义")
    D.p(f"附件1包含100条片段、{len(set(s['video_id'] for s in samples))}个原视频。总时长{result['clip_seconds']:.3f}秒，原文词数{int(summary.words.sum())}。题目提供英文转写；处理保持原文和样本身份不变。")
    D.table(["工具", "输入", "输出/用途"], [
        ["Qwen3 ForcedAligner", "给定转写＋16 kHz音频", "为原文词补充起止时间"],
        ["ModernBERT-base", "完整英文原文", "上下文子词按字符跨度汇总为768维词向量"],
        ["openSMILE eGeMAPSv02", "单声道16 kHz波形", "每个原生声学时间窗25项LLD"],
        ["OpenFace 3.0", "按实际PTS采样的视频帧", "人脸检测、跟踪、28项视觉观测及质量"]], [4.2, 4.3, 8.5])
    D.doc.add_heading("文本768维", 2)
    D.p("ModernBERT在完整原文上冻结推理。利用token字符offset与原词字符跨度对应，去除特殊token，将属于同一词的子词向量取均值。词时间无效时，文本语义仍可保留。")
    D.doc.add_heading("语音50维", 2)
    D.p("openSMILE提取25项LowLevelDescriptors。每个目标区间生成25维加权均值与25维加权总体标准差，拼接为50维。两者使用完全相同的重叠×质量归一化权重，权重之和为1，不作n−1无偏修正，也不除以均值。标准差描述同一区间内源帧的波动；该表示与整段88维Functionals具有不同定义。")
    D.doc.add_heading("视觉56维", 2)
    D.p("单帧28维由8项表情logits、2项视线输出、8项AU原始输出、5个人脸点相对边框坐标10维组成。按目标区间计算加权均值28维与加权总体标准差28维。AU索引不擅自标为FACS编号，预训练表情输出不直接当作赛题标签。")

    D.page("1.1  SAQW-TA：模型与统一表示")
    D.p("将本文方案记为SAQW-TA（Semantic-Anchor Quality-Weighted Temporal Alignment），即语义锚点质量加权时序对齐方法。给定转写的有效词区间提供语义锚点，满足阈值的未分配gap补充目标时间轴，原生音视频观测通过时间重叠与质量加权映射至目标区间。")
    D.image(out / "figures/representation.png", "统一位置表示及三种模态特征空间。")
    D.table(["符号", "定义与实现"], [
        ["Iᵢ、tᵢ", "目标时间区间[sᵢ,eᵢ)及其时间有效标记；对应timestamps与time_valid_mask"],
        ["Tᵢ", "完整原文上下文编码后按字符跨度汇总的词表示；gap无文本；词时间无效不抹去文本"],
        ["Aᵢ、Vᵢ", "对有效原生支持区间计算加权均值与加权总体标准差，公式见第2节"],
        ["Mᵢ、Cᵢ", "三模态0/1可用掩码与逐位置音频/视觉源覆盖；模态缺失与batch填充分别标记"],
        ["Pᵢ", "原词字符/token映射、CSR声学/视觉源索引及权重、原帧号与PTS等来源记录"]], [3, 14])
    D.p("目标区间由词锚点与gap组成，声学窗和视频采样帧保持各自原生时间支持。不同采样率的对应关系由区间重叠确定；该模型未学习情绪表达的跨模态时滞。零源权重或无效词时间使对应音视频掩码为false，不能用补造边界替代。")
    D.doc.add_heading("已实现的方案特点与论证边界", 2)
    D.p("一是词锚点与gap共同保留片段中的时间位置；二是重叠与观测质量共同决定聚合；三是通过来源映射、显式掩码和覆盖集合支持原素材回溯及数值重建。这些特点有代码与产物依据。")
    D.p("方法名用于统一论文叙述；预训练模型及加权统计公式沿用现有工具和数学定义，不因命名宣称首创。是否优于其他对齐方法或提高情感预测效果，需要独立对照证据。")

    D.page("2  共同时间坐标与聚合模型")
    D.p("视频使用ffprobe实际帧PTS，音频使用解码起点偏移加采样下标/16000。共同原点取音视频开始PTS的较早者。原视频帧号、PTS、文本原字符跨度和音频原生窗口始终保留。")
    D.p("Qwen3强制对齐得到词区间Iᵢ=[sᵢ,eᵢ)。源声学帧或视觉支持区间为Jⱼ=[aⱼ,bⱼ)，特征为xⱼ，观测质量为qⱼ。声学有效性使用0/1，视觉使用检测质量并应用已核查的排除规则。")
    D.image(out / "figures/aggregation.png", "时间重叠、质量归一化及区间统计量。")
    D.p("分母为零时，均值与标准差写零，同时将模态掩码置false。NPZ逐位置coverage[i,m]是目标区间Iᵢ与有效源窗口并集的交集长度除以|Iᵢ|；无效词时间记0。它按模态分别计算，采用时间并集去重，与目标时间轴保留率是两个指标。")
    D.table(["时间规则", "具体处理"], [
        ["词区间", "全部原文词均保留；时间无效时不分配音视频来源"],
        ["未分配间隙", "保留≥80 ms的间隙；比较容差1 µs；word_index=-1"],
        ["视觉时间窗", "相邻采样PTS中点形成分段常值支持区间，原帧PTS同时保留"],
        ["名义视觉采样率", "配置约10 Hz，按实际PTS最小间隔抽样；实际频率受原视频帧率影响"],
        ["来源追溯", "每个输出位置保存源索引和归一化权重，可回查原始素材"]], [4.2, 12.8])
    D.p("时间组织应称WORD + GAP。gap仅表示没有分配给结构有效词的区间，可能包含停顿、非语言声音或时间检查失败的词，不能直接称为pause或silence。小于阈值的间隙可能不保留，完整程度由目标时间轴保留率量化。")

    D.page("3  缺失、异常与视觉质量规则")
    D.table(["情况", "处理规则"], [
        ["整段数字静音", "波形峰值≤1e-4：音频全部无效、词时间全部无效；原文保留，视觉独立处理"],
        ["对齐异常", "非正时长、重叠或非单调、越界分别记录；不通过均分时长补造边界"],
        ["无人脸/检测异常", "视觉质量为0，不参与聚合；不能用零向量是否出现来替代掩码"],
        ["人脸跟踪", "初始最大人脸，其后优先最大IoU；未实现说话人身份认证"],
        ["已确认误报", "以视频SHA256绑定原帧排除规则，保留原始检测用于回溯"],
        ["填充", "单样本变长存储；批处理补零并使padding位置与模态掩码为false"]], [4.3, 12.7])
    D.p("视觉阈值0.8，跟踪IoU阈值0.2，输入长边最多960像素。检测分数是观测质量权重，不是情感预测置信度，不能保证人脸未遮挡或就是实际说话人。")
    D.p("鸟类片段 -NFrJFQijFE__2 的原帧45、48、60、63、66、69、96、117，经视觉复核确认检测框落在羽毛上。8帧排除规则已进入正式聚合及预览；源帧与排除记录保留。该复核来自助手视觉检查，不代表全部视频都已完成人工审核。")
    D.p("自动时间有效性仅反映结构约束。对齐可疑样本仍保留在完整输出中；逐词原因见第5.2节。")

    D.page("3.1  观测质量权重的精确定义")
    D.p("视觉观测采用检测得分与有效性门控组成的质量代理量。设cₖ为选中人脸框的有限检测得分，无可用框时取0；dₖ、bₖ、fₖ、rₖ均为0/1标记。下式描述最终参与聚合的有效质量。")
    D.image(out / "figures/quality.png", "视觉质量门控与音频有效性权重。")
    D.table(["项", "代码对应的条件"], [
        ["cₖ", "选中框的检测得分；不是表情分类置信度"],
        ["dₖ", "存在满足检测阈值0.8的选中人脸框"],
        ["bₖ", "框裁到图像范围并转为整数裁剪坐标后，right>left且bottom>top"],
        ["fₖ", "清零之前的28维模型观测均为有限数值，否则整帧质量为0"],
        ["rₖ", "视频SHA256与复核规则一致，且原帧号在排除列表中时为1"],
        ["uₖ、fₖ⁽ᴬ⁾", "音频所属整段非数字静音、当前25维LLD全部有限；二者合取即acoustic_valid"]], [3, 14])
    D.p("实现顺序：VisionEncoder对合格观测保存检测得分，其他情况保存0；reviewed_quality再把复核排除帧置0。原始检测结果保留。聚合器还要求源窗口有限、时长为正，特征与质量有限。")
    D.p("IoU负责目标人脸选择和跟踪，不乘入质量数值，也不提供说话人身份认证。当前没有独立遮挡度、清晰度或说话人置信度项；检测分数尚未校准为这些可靠性概率。音频的整段数字静音规则不是逐帧语音活动检测。")
    D.p("聚合使用重叠长度×q作为权重；coverage只判断是否存在有效支持并对区间取并集，不再按检测得分对时长打折。质量权重与覆盖率承担不同作用。")

    D.page("4  输出接口与来源映射")
    D.p("每条样本保存一份NPZ和一份配套JSON，schema_version=3。时间戳为FP64，正式特征默认FP16，原生源特征缓存为FP32。NPZ读取不需要pickle。")
    D.table(["字段", "形状", "含义"], [
        ["text / audio / vision", "L×768 / L×50 / L×56", "统一目标位置上的三模态特征"],
        ["timestamps", "L×2", "相对共同原点的秒数"],
        ["valid_mask / time_valid_mask", "L / L", "实际位置／时间有效性"],
        ["modality_mask", "L×3", "文本、音频、视觉可用性"],
        ["coverage", "L×2", "逐位置声学、视觉源支持覆盖率"],
        ["word_indices", "L", "原文词索引，间隙为−1"],
        ["source_<m>_intervals", "S×2", "acoustic/vision原生时间窗"],
        ["map_<m>_offsets", "L+1", "CSR行偏移"],
        ["map_<m>_indices / weights", "nnz / nnz", "源行号与归一化权重"],
        ["source_acoustic_valid", "Sₐ", "声学源有效性"],
        ["source_vision_quality / excluded", "Sᵥ / Sᵥ", "复核后视觉质量与排除标记"]], [7.0, 4.3, 5.7], 8.5)
    D.p("位置i的来源为indices[offsets[i]:offsets[i+1]]，权重取相同切片。视觉源行号继续映射到原始frame_index和PTS；文本通过word_indices和字符跨度回到原文。")
    D.p("JSON记录原文、词时间、token映射、原帧号与检测记录；run_config.json集中保存实际配置和环境。cache目录是可重算来源，features目录是正式输出。")
    D.p(f"validation.json的temporal_coverage给出全量覆盖指标，word_alignment记录无效原因，modality_availability区分位置可用率和词级抽取成功率。invalid_words.csv保留{alignment['invalid_words']}个无效词的全部标签及原始边界。")

    D.page("5  全量结果与验证")
    D.table(["检查项", "当前结果"], [
        ["完整样本", f"{result['complete']}/{result['expected']}"],
        ["维度", "每条样本文本768、语音50、视觉56"],
        ["原文词／输出位置", f"{int(summary.words.sum())}／{int(summary.positions.sum())}"],
        ["结构有效词时间", f"{alignment['time_valid_words']}/{alignment['total_words']}，{alignment['structural_valid_ratio']:.4%}"],
        ["目标时间轴保留率", f"{coverage['timeline_retention']['ratio']:.4%}；集合与分母见第5.1节"],
        ["最长未纳入目标时间轴的间隙", f"{1000*result['longest_unmapped_seconds']:.1f} ms（不是最长音频/视觉缺失）"],
        ["全静音／对齐可疑", f"{len(result['silent_audio_ids'])}／{len(result['alignment_suspect_ids'])}条"],
        ["已排除视觉误报", f"{result['review_excluded_vision_frames']}个原始视频帧"],
        ["正式特征及映射大小", f"{result['feature_bytes']/1e6:.3f} MB"]], [7, 10])
    D.doc.add_heading("位置模态可用率与抽取成功率", 2)
    D.table(["指标", "分子／分母", "比例"], [
        ["文本位置可用率", f"{availability['text']['count']}/{availability['positions']}", f"{availability['text']['ratio']:.2%}"],
        ["音频位置可用率", f"{availability['audio']['count']}/{availability['positions']}", f"{availability['audio']['ratio']:.2%}"],
        ["视觉位置可用率", f"{availability['vision']['count']}/{availability['positions']}", f"{availability['vision']['ratio']:.2%}"],
        ["文本词级抽取成功率", f"{availability['text_word_extraction']['count']}/{alignment['total_words']}", f"{availability['text_word_extraction']['ratio']:.2%}"]], [7.4, 5.2, 4.4])
    D.p(f"全部{availability['positions']}个位置包含{alignment['total_words']}个词和{availability['gap_positions']}个gap。gap没有文本，因此T/A/V位置可用率不是抽取成功率，也不是按秒计算的时间覆盖率。例如T=75%表示75%的输出位置有文本，不能据此认定25%的原词抽取失败。")
    D.p("核验输入与缓存哈希、维度、有限值、掩码、完整原文身份、CSR来源和权重及原帧PTS；独立重建加权均值和加权总体标准差，核对导出数值。当前结论限于工程完整性、时间结构与数值可重建性，不提供独立的词边界准确率结论。")

    D.page("5.1  覆盖率：集合、分母与全量结果")
    D.p("对片段s，Dₛ为共同时间坐标下的片段时长，tₛᵢ为time_valid_mask。𝓘ₛ是所有时间有效目标区间的并集，包含词区间及保留gap。𝓥ₛ⁽ᵐ⁾是特征与窗口有限、窗口时长为正、有效质量>0的原生源观测集合；𝓙ₛ⁽ᵐ⁾为其时间窗并集。A代表音频，V代表视觉。")
    D.image(out / "figures/coverage.png", "逐位置coverage、时间轴保留率与四种全量源覆盖率。", 16.4)
    D.p("逐位置Cₛᵢ⁽ᵐ⁾仅对时间有效、正时长目标区间计算，否则存0。值接近1表示有效源支持完整，接近0表示支持稀少或不存在。竖线表示去重后的区间总长度，不能解释为对齐精度。")
    D.table(["指标", "分子/秒", "分母/秒", "比例"], [
        [label, f"{coverage[name]['seconds']:.6f}", f"{coverage[name]['denominator_seconds']:.6f}", f"{coverage[name]['ratio']:.4%}"]
        for name, label in [("timeline_retention", "目标时间轴保留率"), ("audio", "音频有效源覆盖 C_A"),
                            ("vision", "视觉有效源覆盖 C_V"), ("union", "至少一种有效 C_A∪V"),
                            ("intersection", "音视频共同有效 C_A∩V")]], [5.5, 4, 4, 3.5], 8.5)
    D.p(f"{coverage['timeline_retention']['ratio']:.4%}的分母是视频总时长，分子仅取目标区间并集，完全不检查源模态有效性；即使某段音频和视觉都不可用，保留的gap仍可使目标时间轴完整。四种源覆盖率的共同分母是目标区间并集时长，分子均先与𝓘ₛ相交，再在每条片段内去重。全量指标用时长求和后相除，不平均各片段百分比。")
    D.p("源覆盖按有效性取集合，未用检测分数对时长打折。视觉沿用相邻采样PTS中点构造的支持区间。源支持覆盖率不是语音活动占比、逐帧人脸识别精度或词边界准确率；文本是离散词节点，本表不声称三模态共同时间覆盖。")

    D.page("5.2  无效词时间：原因与保留原则")
    D.p(f"{alignment['total_words']}个原文词中，{alignment['time_valid_words']}个通过结构检查（{alignment['structural_valid_ratio']:.4%}），{alignment['invalid_words']}个未通过。以下按“静音→越界→非单调/重叠→非正时长→其他”的展示优先级作互斥归类，每个无效词只计一次；此优先级不是对真实失效根因的推断。")
    reason_rows = [[REASON_LABELS[k], n, f"{n/alignment['total_words']:.2%}", f"{n/alignment['invalid_words']:.2%}"] for k,n in alignment["exclusive_reasons"].items()]
    reason_rows.extend([["未返回时间/未覆盖原文", alignment["unreturned_words_in_validated_samples"], "0.00%", "0.00%"], ["无效词合计", alignment['invalid_words'], f"{alignment['invalid_words']/alignment['total_words']:.2%}", "100.00%"]])
    D.table(["互斥展示分类", "词数", "占全部原词", "占无效词"], reason_rows, [6.5, 2, 4.25, 4.25])
    raw = alignment["nonexclusive_reasons"]
    D.p(f"原始标签允许重复：数字静音{raw['silent_audio']}词、边界越界{raw['out_of_audio_bounds']}词、非正时长{raw['non_positive_duration']}词、非单调/重叠{raw['overlap_or_nonmonotonic']}词。不能将这些标签数直接相加当作独立词数。非正时长包括起止相等或倒置，具体原始时间与所有标签见invalid_words.csv。")
    D.p("全部原文均通过严格字符覆盖核验，当前没有未返回词。若强制对齐没有返回原文的某个词，流程会使该样本失败，不会均分时长或静默丢词。对已有时间但检查不通过的词，保留文本与原始边界，显式标记time_valid=false，不分配音视频来源；相关时间由gap规则决定是否保留。")
    D.p(f"{alignment['samples_with_invalid_words']}条样本至少含一个无效词。报告中的{len(result['alignment_suspect_ids'])}条“对齐可疑”是另一个样本级规则：数字静音或结构有效词比例低于60%。其余样本也可能有少量无效词，因此11条不能解释为全部异常样本数。")

    D.page("6  典型样本的共同时间轴")
    D.p(example["id"] + "\n" + example["raw_text"])
    D.image(typical, "词／gap、波形、声学与AU曲线、逐位置覆盖和带质量q的原始视频帧。")
    words = read_json(example_dir / "align.json")["words"]
    D.table(["原词", "起点/秒", "终点/秒", "时间有效"], [[w["text"], f"{w['start']:.3f}", f"{w['end']:.3f}", str(w["time_valid"])] for w in words[:5]], [5, 4, 4, 4])
    D.p("绿色表示结构有效词时间，橙色表示无效词时间，灰色独立带为未分配gap。coverage曲线来自正式NPZ；帧标题中的q为复核后的观测质量。图中时间为自动结果，可通过CSR来源查询回到原素材。")

    D.page("7  复现、版本与运行记录")
    D.table(["模块", "模型与固定修订号"], [[n, cfg["models"][k]["id"]+"\n"+cfg["models"][k]["revision"]] for k,n in [("align","强制对齐"),("text","文本编码"),("face","视觉编码")]], [3.2, 13.8], 8.8)
    D.table(["依赖", "实际版本"], [[k, record["environment"][k]] for k in ["torch","torchvision","transformers","numpy","opensmile","openface-test","timm"]], [7, 10])
    D.p("关键配置：16 kHz单声道；名义视频采样约10 Hz；检测阈值0.8；跟踪IoU 0.2；长边960；最小间隙80 ms；导出FP16。openSMILE版本与实际字段顺序均记录。")
    D.p("运行：bash q1/install.sh → python -m q1 doctor → python -m q1 run → python -m q1 validate → python -m q1 report。以上python应使用项目的.venv-q1/bin/python。")
    D.p("全部样本由manifest.json统一编号；events.jsonl记录阶段结果；缓存通过输入、配置、代码、依赖和产物哈希校验。当前精简迁移仅复用已验证且提取代码未变的媒体、对齐、文本与视觉结果，语音及聚合已重新执行，详见cleanup.json。")
    D.p("模型说明来源：Qwen/Qwen3-ForcedAligner-0.6B-hf；answerdotai/ModernBERT-base；audeering.github.io/opensmile-python；github.com/CMU-MultiComp-Lab/OpenFace-3.0。题目依据为E题目录中的问题描述文档。")

    for page in range(4):
        D.page(f"附录A  全量100条样本（{page*25+1}—{(page+1)*25}）")
        D.p("每行均包含文本T、音频A、视觉V三模态，维度768/50/56。D为片段时长，L为词＋gap位置数。T/A/V是以全部L位置为分母的位置模态可用率，不是抽取成功率或按秒计算的时间覆盖率；gap不具有文本。")
        rows = []
        for i, (_, r) in enumerate(summary.iloc[page*25:(page+1)*25].iterrows(), page*25+1):
            flags = "/".join((["静"] if r.silent_audio else [])+(["疑"] if r.alignment_suspect else [])+(["无脸"] if r.valid_vision_ratio == 0 else [])) or "—"
            rows.append([i, r.sample_id, f"{r.duration:.3f}", int(r.words), int(r.positions), f"{100*r.valid_text_ratio:.0f}/{100*r.valid_audio_ratio:.0f}/{100*r.valid_vision_ratio:.0f}", flags])
        D.table(["序", "样本ID", "D/秒", "词数", "L", "位置可用T/A/V%", "标记"], rows, [.9, 5.1, 1.7, 1.2, 1.0, 4.0, 3.1], 8)
        D.p("静：整段数字静音；疑：结构性对齐可疑；无脸：当前规则下无有效视觉位置。完整数值和来源映射见对应正式NPZ/JSON。")

    D.page("附录B  25项原生声学指标")
    D.p("顺序直接读取当前openSMILE配置的实际输出。每项在目标区间生成一个加权均值和一个加权总体标准差，共50维；二者采用相同的归一化权重。")
    D.table(["序号", "实际字段名"], [[i+1, name] for i,name in enumerate(names)], [1.5, 15.5], 9)
    D.p("本文档、图表及全量汇总均由当前验证通过的数据生成。更新配置、模型、排除规则或时间映射后，应重新运行受影响阶段并核验，再生成文档。")
    path = out / "问题1_方法流程验证与结果报告.docx"; D.doc.save(path)
    source_names = ["manifest.json", "summary.csv", "validation.json", "run_config.json", "invalid_words.csv"]
    artifact = {"docx": path.name, "docx_sha256": file_hash(path), "figures": D.images, "tables": D.tables,
                "dimensions": result["dimensions"], "schema_version": 3,
                "source_sha256": {name: file_hash(root / name) for name in source_names},
                "generator_sha256": file_hash(__file__), "metrics_version": result["metrics_version"],
                "method": "SAQW-TA: Semantic-Anchor Quality-Weighted Temporal Alignment",
                "validation_code_sha256": file_hash(Path(__file__).with_name("validation.py")),
                "preview_code_sha256": file_hash(Path(__file__).with_name("preview.py"))}
    if shutil.which("libreoffice"):
        with tempfile.TemporaryDirectory(prefix="q1-report-lo-") as profile:
            subprocess.run(["libreoffice", "-env:UserInstallation="+Path(profile).as_uri(), "--headless", "--convert-to", "pdf", "--outdir", str(out), str(path)], check=True, capture_output=True, text=True, timeout=60)
        pdf = path.with_suffix(".pdf")
        if not pdf.exists(): raise RuntimeError("LibreOffice未生成PDF")
        artifact.update(pdf=pdf.name, pdf_sha256=file_hash(pdf))
    atomic_json(out / "manifest.json", artifact)
    return path
